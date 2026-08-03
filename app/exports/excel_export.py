import openpyxl
from openpyxl.styles import Font

from app.calculator.calculator import EstimateCalculation

_HEADERS = [
    "Этап",
    "Работа",
    "Объём",
    "Ед. изм.",
    "Цена работы",
    "Цена материала",
    "Стоимость работы",
    "Стоимость материала",
    "Итого",
]


def export_estimate_to_excel(
    calculation: EstimateCalculation,
    recommendations: list[str],
    output_path: str,
) -> str:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Смета"

    sheet.append(_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for stage in calculation.stages:
        sheet.append(
            [
                stage.key,
                stage.work,
                stage.volume,
                stage.unit,
                stage.work_price if stage.matched else "не найдено в прайсе",
                stage.material_price if stage.matched else "",
                stage.work_cost,
                stage.material_cost,
                stage.total,
            ]
        )

    total_row = ["", "", "", "", "", "", "", "Общий итог", calculation.grand_total]
    sheet.append(total_row)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    if recommendations:
        sheet.append([])
        sheet.append(["Рекомендации (возможно пропущенные этапы):"])
        sheet[sheet.max_row][0].font = Font(bold=True)
        for recommendation in recommendations:
            sheet.append([recommendation])

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 50)

    workbook.save(output_path)
    return output_path
