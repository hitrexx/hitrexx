import os
from dataclasses import dataclass, field

import docx
import fitz  # PyMuPDF

_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


@dataclass
class ProjectContent:
    text: str = ""
    image_paths: list[str] = field(default_factory=list)


def _extract_pdf(file_path: str, render_dir: str) -> tuple[str, list[str]]:
    os.makedirs(render_dir, exist_ok=True)
    text_parts: list[str] = []
    image_paths: list[str] = []

    document = fitz.open(file_path)
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    try:
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text().strip()
            if page_text:
                text_parts.append(f"[Страница {page_number}]\n{page_text}")

            pixmap = page.get_pixmap(dpi=150)
            image_path = os.path.join(render_dir, f"{base_name}_p{page_number}.png")
            pixmap.save(image_path)
            image_paths.append(image_path)
    finally:
        document.close()

    return "\n\n".join(text_parts), image_paths


def _extract_docx(file_path: str) -> str:
    document = docx.Document(file_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))

    return "\n".join(paragraphs + table_lines)


def collect_project_content(file_paths: list[str], render_dir: str) -> ProjectContent:
    content = ProjectContent()

    for file_path in file_paths:
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            text, images = _extract_pdf(file_path, render_dir)
            if text:
                content.text += f"\n\n=== {os.path.basename(file_path)} ===\n{text}"
            content.image_paths.extend(images)
        elif ext == ".docx":
            text = _extract_docx(file_path)
            if text:
                content.text += f"\n\n=== {os.path.basename(file_path)} ===\n{text}"
        elif ext in _IMAGE_EXTENSIONS:
            content.image_paths.append(file_path)

    return content
