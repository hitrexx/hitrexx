import openpyxl
import pytest

from app.parser.price_list_parser import PriceListError, parse_price_list


def _make_price_list(tmp_path, rows):
    path = tmp_path / "price_list.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Наименование", "Единица", "Цена работы", "Цена материала"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return str(path)


def test_parse_price_list_happy_path(tmp_path):
    path = _make_price_list(
        tmp_path,
        [
            ["101", "Устройство ленточного фундамента", "м3", 1500, 3200],
            ["102", "Кладка стен из газобетона", "м2", 900, 2100],
        ],
    )

    items = parse_price_list(path)

    assert len(items) == 2
    assert items[0]["name"] == "Устройство ленточного фундамента"
    assert items[0]["unit"] == "м3"
    assert items[0]["work_price"] == 1500.0
    assert items[0]["material_price"] == 3200.0


def test_parse_price_list_skips_empty_rows(tmp_path):
    path = _make_price_list(
        tmp_path,
        [
            ["101", "Устройство фундамента", "м3", 1500, 3200],
            [None, None, None, None, None],
        ],
    )

    items = parse_price_list(path)
    assert len(items) == 1


def test_parse_price_list_missing_columns_raises(tmp_path):
    path = tmp_path / "bad.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Что-то", "Другое"])
    sheet.append(["a", "b"])
    workbook.save(path)

    with pytest.raises(PriceListError):
        parse_price_list(str(path))


def test_parse_price_list_finds_header_below_title_rows(tmp_path):
    """Real sметчик calculators often have a few title/summary rows before
    the actual table header, and use a single combined price column."""
    path = tmp_path / "real_world.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Площадь дома,м2", None, 0])
    sheet.append(["ФИО заказчика:", None, None])
    sheet.append(["№ п.п", "Наименование работ, материалов, затрат", "Единица изм.", "Цена за единицу", "Количество"])
    sheet.append(["Раздел 1. Фундамент"])
    sheet.append([1, "Арматура рифленая D-12мм, м", "м", 65, 320])
    sheet.append([2, "Бетон М250 с доставкой, м3", "м3", 9300, 22])
    workbook.save(path)

    items = parse_price_list(str(path))

    assert len(items) == 2
    assert items[0]["name"] == "Арматура рифленая D-12мм, м"
    assert items[0]["unit"] == "м"
    assert items[0]["work_price"] == 65.0
    assert items[0]["material_price"] == 0.0


def test_parse_price_list_uses_first_sheet_with_a_valid_table(tmp_path):
    path = tmp_path / "multi_sheet.xlsx"
    workbook = openpyxl.Workbook()
    reference_sheet = workbook.active
    reference_sheet.title = "Справочник"
    reference_sheet.append(["Просто текст без таблицы"])

    price_sheet = workbook.create_sheet("Прайс")
    price_sheet.append(["Наименование", "Единица", "Цена работы", "Цена материала"])
    price_sheet.append(["Кладка стен", "м2", 900, 2100])
    workbook.save(path)

    items = parse_price_list(str(path))

    assert len(items) == 1
    assert items[0]["name"] == "Кладка стен"
